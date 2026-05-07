"""Export case data to Word (.docx) format."""
import os
import logging
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

logger = logging.getLogger(__name__)


def set_font(run, name="微软雅黑", size=10.5):
    """Set font for a run (supports Chinese)."""
    run.font.name = name
    run.font.size = Pt(size)
    run.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_heading(doc, text, level=1):
    """Add heading with Chinese font support."""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return h


def add_bullet(doc, text):
    """Add bullet point."""
    p = doc.add_paragraph(text, style="List Bullet")
    for run in p.runs:
        set_font(run)
    return p


def add_info_item(doc, label, value):
    """Add a labeled info item (bold label + normal value)."""
    p = doc.add_paragraph()
    run_label = p.add_run(f"{label}：")
    run_label.bold = True
    set_font(run_label)
    run_value = p.add_run(value or "未提及")
    set_font(run_value)
    return p


def generate_docx(case_data: dict, output_path: str) -> str:
    """Generate a Word document from case data."""
    doc = Document()
    
    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "微软雅黑"
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    
    # Title
    title = doc.add_heading(case_data.get("title", "排障案例"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.name = "微软雅黑"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    
    # Basic Info
    add_heading(doc, "基本信息", level=1)
    add_info_item(doc, "故障现象", case_data.get("fault_phenomenon"))
    add_info_item(doc, "涉及设备/软件", case_data.get("equipment"))
    add_info_item(doc, "故障时间", case_data.get("fault_time"))
    add_info_item(doc, "处理人员", case_data.get("personnel"))
    
    # Fault Description
    add_heading(doc, "故障描述", level=1)
    p = doc.add_paragraph(case_data.get("fault_description", "未提及"))
    for run in p.runs:
        set_font(run)
    
    # Troubleshooting Steps
    add_heading(doc, "排障过程", level=1)
    steps = case_data.get("troubleshooting_steps", [])
    if steps:
        for step in steps:
            add_heading(doc, f"步骤 {step.get('step', '?')}", level=2)
            add_info_item(doc, "操作", step.get("action"))
            add_info_item(doc, "观察结果", step.get("observation"))
            add_info_item(doc, "分析判断", step.get("analysis"))
    else:
        doc.add_paragraph("未提及排查步骤")
    
    # Root Cause
    add_heading(doc, "根因分析", level=1)
    p = doc.add_paragraph(case_data.get("root_cause", "未提及"))
    for run in p.runs:
        set_font(run)
    
    # Solution
    add_heading(doc, "解决方案", level=1)
    p = doc.add_paragraph(case_data.get("solution", "未提及"))
    for run in p.runs:
        set_font(run)
    
    # Lessons Learned
    add_heading(doc, "经验总结", level=1)
    lessons = case_data.get("lessons_learned", [])
    if lessons:
        for lesson in lessons:
            add_bullet(doc, lesson)
    else:
        doc.add_paragraph("未提及经验总结")
    
    doc.save(output_path)
    logger.info(f"Word document saved to {output_path}")
    return output_path
